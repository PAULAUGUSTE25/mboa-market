"""
Analyze the ORIGINAL document to understand:
1. Where are the images?
2. What text is BEFORE each image?
3. What text is AFTER each image (Explanation)?
4. What is the figure caption?

Then we can properly transform the explanations.
"""

from docx import Document
from docx.oxml.ns import qn
import os

INPUT_FILE = r"c:\Users\HP\Desktop\mboa-market\NDEO PAUL AUGUSTE ICTU20212527_COMPLET.docx"

def analyze_document():
    """Analyze the document structure around images"""
    print("="*80)
    print("ANALYZING ORIGINAL DOCUMENT STRUCTURE")
    print("="*80)
    
    doc = Document(INPUT_FILE)
    
    # Find all paragraphs and identify images
    image_locations = []
    
    for i, para in enumerate(doc.paragraphs):
        # Check if paragraph contains an image
        has_image = False
        for run in para.runs:
            if run._element.xpath('.//a:blip'):
                has_image = True
                break
            # Also check for inline shapes
            if run._element.xpath('.//w:drawing'):
                has_image = True
                break
        
        if has_image:
            image_locations.append(i)
    
    print(f"\nFound {len(image_locations)} paragraphs with images")
    print("\n" + "="*80)
    
    # For each image, show context
    for img_idx, para_idx in enumerate(image_locations):
        print(f"\n{'='*80}")
        print(f"IMAGE #{img_idx + 1} at paragraph {para_idx}")
        print("="*80)
        
        # Show 3 paragraphs BEFORE
        print("\n--- BEFORE IMAGE (3 paragraphs) ---")
        for j in range(max(0, para_idx - 3), para_idx):
            text = doc.paragraphs[j].text.strip()
            if text:
                print(f"  [{j}] {text[:150]}{'...' if len(text) > 150 else ''}")
        
        # Show the image paragraph
        print(f"\n--- IMAGE PARAGRAPH [{para_idx}] ---")
        img_para_text = doc.paragraphs[para_idx].text.strip()
        if img_para_text:
            print(f"  Caption/Text: {img_para_text}")
        else:
            print("  [Image only, no text]")
        
        # Show 3 paragraphs AFTER
        print("\n--- AFTER IMAGE (3 paragraphs) ---")
        for j in range(para_idx + 1, min(len(doc.paragraphs), para_idx + 4)):
            text = doc.paragraphs[j].text.strip()
            if text:
                print(f"  [{j}] {text[:150]}{'...' if len(text) > 150 else ''}")
        
        if img_idx >= 10:  # Limit output
            print(f"\n... and {len(image_locations) - 11} more images")
            break
    
    return image_locations


def export_full_analysis():
    """Export complete analysis to a text file"""
    print("\n\n" + "="*80)
    print("EXPORTING FULL ANALYSIS TO FILE")
    print("="*80)
    
    doc = Document(INPUT_FILE)
    output_file = r"c:\Users\HP\Desktop\mboa-market\DOCUMENT_ANALYSIS.txt"
    
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write("COMPLETE DOCUMENT ANALYSIS\n")
        f.write("="*80 + "\n\n")
        
        image_count = 0
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            
            # Check for image
            has_image = False
            for run in para.runs:
                if run._element.xpath('.//a:blip') or run._element.xpath('.//w:drawing'):
                    has_image = True
                    break
            
            if has_image:
                image_count += 1
                f.write(f"\n{'='*80}\n")
                f.write(f"[PARAGRAPH {i}] === IMAGE #{image_count} ===\n")
                f.write(f"{'='*80}\n")
                if text:
                    f.write(f"Caption: {text}\n")
            elif text:
                # Check if this looks like an explanation
                text_lower = text.lower()
                is_explanation = any(pattern in text_lower for pattern in [
                    "this diagram", "this figure", "this table", "this image",
                    "this screenshot", "this flowchart", "this curve",
                    "the diagram", "the figure", "the table",
                    "as shown", "as illustrated", "we can see", "we can observe",
                    "presents the", "shows the", "illustrates the",
                    "explanation", "below is", "above is"
                ])
                
                if is_explanation:
                    f.write(f"\n[PARAGRAPH {i}] **EXPLANATION**\n")
                    f.write(f"{text}\n")
                else:
                    f.write(f"\n[PARAGRAPH {i}]\n")
                    f.write(f"{text}\n")
        
        f.write(f"\n\n{'='*80}\n")
        f.write(f"TOTAL IMAGES FOUND: {image_count}\n")
        f.write(f"{'='*80}\n")
    
    print(f"\n✅ Full analysis exported to: {output_file}")
    print(f"   Total images found: {image_count}")
    print("\nPlease review this file to see the exact structure of your document.")


if __name__ == "__main__":
    image_locations = analyze_document()
    export_full_analysis()
